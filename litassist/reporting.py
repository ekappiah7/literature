"""Reporting: PRISMA-ScR flow diagram, search methods text and a grounded draft.

The methods paragraph is written from the search log, not by the AI, so it is
exactly what was done. The AI draft may only cite records included in the
review; every citation marker is checked and the reference list is built from
the stored records, never from the AI's text.
"""

import io
import re
from datetime import datetime

import matplotlib

matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.shared import Pt  # noqa: E402

from litassist.ai import AI, AIError, protocol_text  # noqa: E402


# PRISMA-ScR flow diagram

def prisma_png(c: dict) -> bytes:
    fig, ax = plt.subplots(figsize=(8.27, 9.5), dpi=200)
    ax.set_xlim(0, 100)
    ax.set_ylim(0, 100)
    ax.axis("off")
    box = dict(boxstyle="square,pad=0.6", facecolor="white", edgecolor="#1F4E79", linewidth=1.2)
    side = dict(boxstyle="square,pad=0.6", facecolor="#F2F6FA", edgecolor="#1F4E79", linewidth=1.0)

    sources = "\n".join(f"{name}: n = {n}" for name, n in c["by_source"].items()) or "No searches yet"
    ft_reasons = "\n".join(f"{r}: n = {n}" for r, n in c["ft_excluded_reasons"].items()) or "None yet"
    rows = [
        (88, f"Records identified from databases\nand other sources (n = {c['records_identified']})\n{sources}",
         f"Duplicate records removed\n(n = {c['duplicates_removed']})"),
        (66, f"Records screened on title and abstract\n(n = {c['records_screened']})",
         f"Records excluded (n = {c['excluded']})\nNot yet screened (n = {c['undecided']})"),
        (46, f"Full texts sought for retrieval\n(n = {c['ft_sought']})",
         f"Full texts not retrieved\n(n = {c['ft_not_retrieved']})"),
        (27, f"Full texts assessed for eligibility\n(n = {c['ft_assessed']})",
         f"Full texts excluded (n = {c['ft_excluded']})\n{ft_reasons}"),
        (8, f"Sources of evidence included\nin the review (n = {c['ft_included']})", ""),
    ]
    for i, (y, main, right) in enumerate(rows):
        ax.text(30, y, main, ha="center", va="center", fontsize=8.5, bbox=box, wrap=True)
        if right:
            ax.text(80, y, right, ha="center", va="center", fontsize=8, bbox=side)
            ax.annotate("", xy=(64, y), xytext=(47, y), arrowprops=dict(arrowstyle="->", color="#1F4E79"))
        if i < len(rows) - 1:
            ax.annotate("", xy=(30, rows[i + 1][0] + 5.5), xytext=(30, y - 6),
                        arrowprops=dict(arrowstyle="->", color="#1F4E79"))
    ax.text(2, 98, "Identification", fontsize=9, weight="bold", color="#1F4E79")
    ax.text(2, 74, "Screening", fontsize=9, weight="bold", color="#1F4E79")
    ax.text(2, 15, "Included", fontsize=9, weight="bold", color="#1F4E79")
    buf = io.BytesIO()
    fig.savefig(buf, format="png", bbox_inches="tight")
    plt.close(fig)
    return buf.getvalue()


# Methods text from the search log

def _day(stamp: str) -> str:
    try:
        return datetime.strptime(stamp[:10], "%Y-%m-%d").strftime("%d %B %Y").lstrip("0")
    except ValueError:
        return stamp[:10]


def methods_text(searches: list[dict], c: dict) -> str:
    if not searches:
        return "No searches have been run yet."
    db_runs = [s for s in searches if not s["source"].startswith(("Imported", "Citation"))]
    imports = [s for s in searches if s["source"].startswith("Imported")]
    chasing = [s for s in searches if s["source"].startswith("Citation")]
    parts = []
    latest = {}
    for s in db_runs:
        latest[s["source"]] = s
    if latest:
        names = ", ".join(sorted(latest))
        last_date = _day(max(s["run_at"] for s in latest.values()))
        limits = {s["filters"] for s in latest.values() if s["filters"]}
        parts.append(
            f"We searched {names}, with the most recent search on {last_date}. "
            + (f"Searches were limited by publication date ({'; '.join(sorted(limits))}). " if limits
               else "No date or language limits were applied. ")
            + "The full search strings are given in the supplementary material."
        )
    if imports:
        parts.append(f"Records were also imported from {len(imports)} file export(s) "
                     f"({', '.join(sorted({s['source'].replace('Imported: ', '') for s in imports}))}).")
    if chasing:
        parts.append(f"Reference lists and citing articles of included sources were checked "
                     f"through OpenAlex ({sum(s['retrieved'] for s in chasing)} records retrieved).")
    parts.append(
        f"In total {c['records_identified']} records were identified and {c['duplicates_removed']} duplicates "
        f"removed, leaving {c['records_screened']} records for title and abstract screening. "
        f"Of these, {c['excluded']} were excluded and {c['ft_sought']} were sought for full text review. "
        f"{c['ft_assessed']} full texts were assessed, {c['ft_excluded']} were excluded, and "
        f"{c['ft_included']} sources of evidence were included."
    )
    return " ".join(parts)


def search_strings_table(searches: list[dict]) -> list[tuple]:
    return [(s["source"], _day(s["run_at"]), s["filters"] or "None", s["hits"], s["query"])
            for s in searches if not s["source"].startswith(("Imported", "Citation"))]


# Citations

def cite_label(r: dict) -> str:
    authors = r["authors"]
    if not authors:
        name = (r["journal"] or "Anonymous").split(":")[0]
    elif len(authors) == 1:
        name = authors[0]["last"]
    elif len(authors) == 2:
        name = f"{authors[0]['last']} and {authors[1]['last']}"
    else:
        name = f"{authors[0]['last']} et al."
    return f"{name}, {r['year'] or 'n.d.'}"


def reference(r: dict) -> str:
    names = [f"{a['last']} {a['initials'] or (a['first'][:1] if a['first'] else '')}".strip() for a in r["authors"]]
    if len(names) > 6:
        names = names[:6] + ["et al."]
    author_text = ", ".join(names) if names else "Anonymous"
    bits = [f"{author_text} ({r['year'] or 'n.d.'}). {r['title'].rstrip('.')}."]
    journal = r["journal"]
    if journal:
        vol = f", {r['volume']}" if r["volume"] else ""
        iss = f"({r['issue']})" if r["issue"] else ""
        pages = f", {r['pages']}" if r["pages"] else ""
        bits.append(f"{journal}{vol}{iss}{pages}.")
    if r["doi"]:
        bits.append(f"https://doi.org/{r['doi']}")
    elif r["pmid"]:
        bits.append(f"PMID: {r['pmid']}")
    return " ".join(bits)


# Grounded draft

DRAFT_SCHEMA = {
    "type": "object",
    "properties": {
        "sections": {
            "type": "array",
            "items": {
                "type": "object",
                "properties": {
                    "heading": {"type": "string"},
                    "paragraphs": {"type": "array", "items": {"type": "string"}},
                },
                "required": ["heading", "paragraphs"],
                "additionalProperties": False,
            },
        }
    },
    "required": ["sections"],
    "additionalProperties": False,
}

DRAFT_SYSTEM = """You write a first draft of the results synthesis for a review, for the researcher to rewrite.

Rules:
- Use only the charted data provided. Do not add facts, numbers or studies from outside knowledge.
- Cite sources only with their markers exactly as given, for example [R12]. Every factual sentence about a study needs a marker.
- Never invent a marker. If the data do not support a point, leave it out.
- Organise the synthesis thematically around the review question, not study by study. Describe patterns, gaps and inconsistencies in reporting.
- Plain, formal academic English. No bullet points. Short paragraphs.
- Answer only with JSON: {{"sections": [{{"heading": ..., "paragraphs": [...]}}]}}

Review protocol:
{protocol}"""


def evidence_block(records: list[dict]) -> str:
    lines = []
    for r in records:
        lines.append(f"[R{r['id']}] {cite_label(r)}: {r['title']}")
        for field, item in (r.get("chart") or {}).items():
            value = item.get("value", "") if isinstance(item, dict) else str(item)
            if value and value != "Not reported":
                lines.append(f"  {field}: {value}")
        if not r.get("chart"):
            lines.append(f"  Abstract: {r['abstract'][:1500]}")
        lines.append("")
    return "\n".join(lines)


MARKER = re.compile(r"\[R(\d+)\]")


def draft(ai: AI, project: dict, records: list[dict], focus: str = "") -> tuple[list[dict], dict, str]:
    """Returns (sections with markers replaced, check report, model)."""
    if not records:
        raise AIError("No sources are included in the review yet.")
    user = (f"{('Focus of this draft: ' + focus) if focus else ''}\n\nCharted data for the included sources:\n\n"
            f"{evidence_block(records)}")
    data, model = ai.json(DRAFT_SYSTEM.format(protocol=protocol_text(project)), user, DRAFT_SCHEMA, max_tokens=16000)
    by_id = {r["id"]: r for r in records}
    used, unknown = set(), set()

    def replace(match):
        rid = int(match.group(1))
        if rid in by_id:
            used.add(rid)
            return f"({cite_label(by_id[rid])})"
        unknown.add(rid)
        return "[citation removed: not an included source]"

    sections = []
    for sec in data.get("sections", []):
        paras = [MARKER.sub(replace, p).replace(") (", "; ") for p in sec.get("paragraphs", []) if p.strip()]
        sections.append({"heading": sec.get("heading", "").strip(), "paragraphs": paras})
    report = {"cited": sorted(used), "invalid_markers": sorted(unknown),
              "not_cited": sorted(set(by_id) - used)}
    return sections, report, model


def draft_docx(project: dict, sections: list[dict], cited: list[dict], methods: str, report: dict, model: str) -> bytes:
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    doc.add_heading(f"{project['name']}: first draft", level=1)
    note = doc.add_paragraph()
    note.add_run(
        f"Generated {datetime.now():%d %B %Y} by LitAssist from charted data using {model}. This is a first draft "
        "to rewrite in your own words. Check every sentence against its source before use."
    ).italic = True
    if report["invalid_markers"]:
        doc.add_paragraph(f"Warning: {len(report['invalid_markers'])} citation(s) pointed to records that are not "
                          "included sources and were removed.")
    doc.add_heading("Search methods", level=2)
    doc.add_paragraph(methods)
    for sec in sections:
        if sec["heading"]:
            doc.add_heading(sec["heading"], level=2)
        for para in sec["paragraphs"]:
            doc.add_paragraph(para)
    doc.add_heading("References", level=2)
    for r in sorted(cited, key=lambda x: (x["authors"][0]["last"] if x["authors"] else "", x["year"])):
        doc.add_paragraph(reference(r))
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()
