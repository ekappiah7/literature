from pathlib import Path

import pytest
from openpyxl import load_workbook
import io

from litassist import db
from litassist.export import to_bibtex, to_excel, to_ris
from litassist.pubmed import parse_pubmed_xml
from litassist.zotero import to_zotero_item

FIXTURE = Path(__file__).parent / "pubmed_sample.xml"


@pytest.fixture
def conn(tmp_path):
    return db.connect(tmp_path / "test.db")


@pytest.fixture
def records():
    return parse_pubmed_xml(FIXTURE.read_bytes())


def test_parse_pubmed(records):
    assert len(records) == 2
    r = records[0]
    assert r["pmid"] == "11111111"
    assert r["doi"] == "10.1000/example.1"
    assert r["pmcid"] == "PMC1234567"
    assert r["title"] == "Live birth after IVF in Kumasi, Ghana: a cohort study."
    assert r["abstract"].startswith("BACKGROUND: Few data")
    assert "RESULTS: Live birth" in r["abstract"]
    assert r["authors"][0] == {"last": "Mensah", "first": "Kofi", "initials": "K"}
    assert r["year"] == "2021"
    assert r["pages"] == "101-110"
    assert "Fertilization in Vitro" in r["mesh"]
    # second record: DOI from ELocationID, MedlineDate year, collective author
    s = records[1]
    assert s["doi"] == "10.1000/example.2"
    assert s["year"] == "2019"
    assert s["authors"][0]["last"] == "Ghana ART Study Group"


def test_dedup_and_decisions(conn, records):
    pid = db.create_project(conn, "Test")
    sid = db.log_search(conn, pid, "PubMed", "ivf", 2)
    assert db.add_records(conn, pid, sid, records) == (2, 0)
    db.finish_search(conn, sid, 2, 2, 0)
    # same records again, one with DOI written as a URL, are duplicates
    again = [dict(records[0], pmid=""), dict(records[1], doi="https://doi.org/10.1000/EXAMPLE.2", pmid="")]
    sid2 = db.log_search(conn, pid, "PubMed", "ivf", 2)
    assert db.add_records(conn, pid, sid2, again) == (0, 2)
    db.finish_search(conn, sid2, 2, 0, 2)

    rec_ids = [r["id"] for r in db.list_records(conn, pid)]
    db.set_decision(conn, rec_ids[0], "include")
    db.set_decision(conn, rec_ids[1], "exclude", "Not sub-Saharan Africa")
    with pytest.raises(ValueError):
        db.set_decision(conn, rec_ids[1], "perhaps")
    c = db.counts(conn, pid)
    expected = {"records_identified": 4, "duplicates_removed": 2, "records_screened": 2,
                "included": 1, "excluded": 1, "maybe": 0, "undecided": 0, "ft_sought": 1}
    assert {k: c[k] for k in expected} == expected
    assert len(db.list_records(conn, pid, "include")) == 1
    log = conn.execute("SELECT COUNT(*) FROM decision_log").fetchone()[0]
    assert log == 2

    # projects are isolated from each other
    other = db.create_project(conn, "Other")
    assert db.add_records(conn, other, None, records) == (2, 0)


def test_exports(conn, records):
    pid = db.create_project(conn, "Export test", question="Q?")
    db.add_records(conn, pid, None, records)
    recs = db.list_records(conn, pid)

    ris = to_ris(recs)
    assert ris.count("TY  - JOUR") == 2
    assert "AU  - Mensah, Kofi" in ris
    assert "SP  - 101" in ris and "EP  - 110" in ris
    assert "DO  - 10.1000/example.1" in ris

    bib = to_bibtex(recs)
    assert bib.count("@article{") == 2
    assert "mensah2021live" in bib
    assert "pages = {101--110}" in bib

    xlsx = to_excel(db.get_project(conn, pid), recs, db.list_searches(conn, pid), db.counts(conn, pid))
    wb = load_workbook(io.BytesIO(xlsx))
    assert wb.sheetnames == ["Records", "Search log", "PRISMA counts", "Protocol"]
    assert wb["Records"].max_row == 3

    item = to_zotero_item(recs[0], "ABCD1234")
    assert item["collections"] == ["ABCD1234"]
    assert item["creators"][0] == {"creatorType": "author", "lastName": "Mensah", "firstName": "Kofi"}
    assert "PMID: 11111111" in item["extra"]
    group = to_zotero_item(recs[1], "ABCD1234")
    assert group["creators"][0] == {"creatorType": "author", "name": "Ghana ART Study Group"}
